import os
import sys
import json
import argparse
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# Import dictionary
try:
    from hs_code_dict import HS_CODES, CUSTOMS_DICTIONARY
except ImportError:
    from .hs_code_dict import HS_CODES, CUSTOMS_DICTIONARY

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pymupdf

def match_hs_code(name_desc):
    text = name_desc.lower()
    if any(k in text for k in ["spot", "proiettore", "track spot"]):
        return HS_CODES["TRACK_SPOT"]
    elif any(k in text for k in ["trimless", "incasso", "downlight"]):
        return HS_CODES["DOWNLIGHT_TRIMLESS"]
    elif any(k in text for k in ["track", "binario", "șină", "sina"]):
        if any(k in text for k in ["feed", "connector", "joint", "cap", "tappo", "giunto", "alimentare"]):
            return HS_CODES["TRACK_ACCESSORY"]
        return HS_CODES["MAGNETIC_TRACK"]
    elif any(k in text for k in ["strip", "striscia", "banda led", "лента"]):
        return HS_CODES["LED_STRIP"]
    elif any(k in text for k in ["driver", "alimentatore", "power supply", "sursa", "блок"]):
        return HS_CODES["POWER_SUPPLY"]
    elif any(k in text for k in ["profil", "profile", "profilo"]):
        return HS_CODES["ALUMINUM_PROFILE"]
    elif any(k in text for k in ["sospensione", "pendant", "suspendat"]):
        return HS_CODES["SUSPENSION_LUMINAIRE"]
    elif any(k in text for k in ["applique", "parete", "wall lamp", "bra"]):
        return HS_CODES["WALL_SCONCE"]
    elif any(k in text for k in ["outdoor", "esterno", "proiettore esterno", "ip65", "ip66"]):
        return HS_CODES["OUTDOOR_LUMINAIRE"]
    return HS_CODES["LED_LUMINAIRE"]

def translate_to_ro(name_desc):
    orig = name_desc
    lower = orig.lower()
    for src, dst in CUSTOMS_DICTIONARY.items():
        if src in lower:
            idx = lower.find(src)
            matched = orig[idx:idx+len(src)]
            orig = orig.replace(matched, dst)
            lower = orig.lower()
    return orig

def generate_customs_excel(data, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Specificație Vamală"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    navy_dark = "1B2A4A"
    border_gray = "CBD5E1"
    thin_border = Border(
        left=Side(style='thin', color=border_gray),
        right=Side(style='thin', color=border_gray),
        top=Side(style='thin', color=border_gray),
        bottom=Side(style='thin', color=border_gray)
    )

    # Title
    ws.merge_cells("A1:J1")
    ws["A1"] = "ANEXĂ LA FACTURĂ - SPECIFICAȚIE VAMALĂ DE MĂRFURI"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Metadata
    supplier = data.get("supplier", "Lombardo S.p.a. (Villongo, BG, Italia)")
    importer = data.get("importer", "Smart Light S.R.L. (Chișinău, str. Pajurii 11, IDNO: 1018600025412)")
    cmr_no = data.get("cmr_no", "CMR MD-IT 89412")
    truck_no = data.get("truck_no", "C LK 421 / TR 890")
    inv_no = data.get("invoice_no", "PRO-2026/09-89")
    inv_date = data.get("invoice_date", "29.09.2026")

    meta_rows = [
        ("Furnizor / Exportator:", supplier, "Nr. Factură Proformă:", inv_no),
        ("Importator / Destinatar:", importer, "Data emiterii:", inv_date),
        ("Document de transport:", cmr_no, "Auto / Remorcă:", truck_no),
        ("Condiție de livrare:", data.get("incoterms", "EXW Villongo (Incoterms 2020)"), "Valută de decontare:", data.get("currency", "EUR"))
    ]

    r = 3
    for label1, val1, label2, val2 in meta_rows:
        ws.cell(row=r, column=1, value=label1).font = Font(name="Arial", size=9, bold=True, color="475569")
        ws.cell(row=r, column=2, value=val1).font = Font(name="Arial", size=9)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(row=r, column=7, value=label2).font = Font(name="Arial", size=9, bold=True, color="475569")
        ws.cell(row=r, column=8, value=val2).font = Font(name="Arial", size=9, bold=True)
        ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=10)
        ws.row_dimensions[r].height = 20
        r += 1

    # Headers
    headers = [
        "Nr.\ncrt.", "Cod Articol\n(Producător)", "Denumirea Mărfii conform Vămii\n(Traducere Oficială Română)",
        "Cod Tarifar Vamal\n(HS Code)", "Țara de\nOrigine", "Cantitate\n(Buc/M)", "Greutate\nNetă (kg)",
        "Greutate\nBrută (kg)", "Preț Unitar\n(EUR)", "Valoare Totală\n(EUR)"
    ]

    header_row = 8
    ws.row_dimensions[header_row].height = 32
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Populate Items
    items = data.get("items", [])
    row_idx = 9
    total_qty = 0
    total_net = 0.0
    total_gross = 0.0
    total_val = 0.0

    for idx, it in enumerate(items, 1):
        code = it.get("code") or it.get("articol") or f"ART-{idx}"
        raw_name = it.get("name") or it.get("denumire") or "Corp iluminat"
        ro_name = translate_to_ro(raw_name)
        hs_info = match_hs_code(raw_name)
        hs_code = it.get("hs_code", hs_info["code"])
        country = it.get("origin_country", "IT")
        qty = float(it.get("qty", 1))
        net_kg = float(it.get("net_kg", round(qty * 0.85, 2)))
        gross_kg = float(it.get("gross_kg", round(net_kg * 1.12, 2)))
        unit_price = float(it.get("price_eur", it.get("unit_price", 50.0)))
        total_p = round(qty * unit_price, 2)

        total_qty += qty
        total_net += net_kg
        total_gross += gross_kg
        total_val += total_p

        row_vals = [
            idx, code, f"{ro_name} ({raw_name})", hs_code, country,
            qty, net_kg, gross_kg, unit_price, total_p
        ]

        for col_idx, val in enumerate(row_vals, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = Font(name="Arial", size=9)
            c.border = thin_border
            if col_idx in [1, 4, 5]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [6, 7, 8]:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = "#,##0.00" if col_idx > 6 else "#,##0"
            elif col_idx in [9, 10]:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = "€#,##0.00"
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

    # Totals Row
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
    ws.cell(row=row_idx, column=1, value="TOTAL GENERAL MARFĂ:").font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="right", vertical="center")

    for col, val, fmt in [(6, total_qty, "#,##0"), (7, total_net, "#,##0.00"), (8, total_gross, "#,##0.00"), (10, total_val, "€#,##0.00")]:
        c = ws.cell(row=row_idx, column=col, value=val)
        c.font = Font(name="Arial", size=10, bold=True)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, 11):
        ws.cell(row=row_idx, column=col).border = thin_border
        ws.cell(row=row_idx, column=col).fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    ws.row_dimensions[row_idx].height = 28

    # Signatures
    r = row_idx + 2
    ws.cell(row=r, column=2, value="Expeditor / Furnizor:").font = Font(name="Arial", size=9, bold=True)
    ws.cell(row=r, column=7, value="Destinatar / Importator:").font = Font(name="Arial", size=9, bold=True)
    r += 1
    ws.cell(row=r, column=2, value="Semnătura și Ștampila: __________________").font = Font(name="Arial", size=9)
    ws.cell(row=r, column=7, value="Director General Smart Light SRL: __________________").font = Font(name="Arial", size=9)

    col_widths = {1: 6, 2: 18, 3: 40, 4: 16, 5: 9, 6: 12, 7: 14, 8: 14, 9: 14, 10: 16}
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path

def generate_proforma_pdf(data, out_path):
    doc = pymupdf.open()
    font_path = "C:/Windows/Fonts/arial.ttf" if os.path.exists("C:/Windows/Fonts/arial.ttf") else None
    
    page_w, page_h = 595.3, 841.9
    margin = 30.0
    page = doc.new_page(width=page_w, height=page_h)

    # 1. Header
    page.draw_rect(pymupdf.Rect(margin, 25, margin + 120, 50), color=(0.11, 0.16, 0.29), fill=(0.11, 0.16, 0.29))
    page.insert_text(pymupdf.Point(margin + 8, 41), "smart | light", fontsize=13, color=(1, 1, 1), fontfile=font_path)
    page.insert_text(pymupdf.Point(margin + 8, 47), "customs logistics hub", fontsize=5.5, color=(0.8, 0.85, 0.95), fontfile=font_path)

    mgr_text = "SMART LIGHT S.R.L.\nIDNO: 1018600025412 | TVA: 0508912\nChisinau, str. Pajurii 11 | sales@smartlight.md"
    page.insert_text(pymupdf.Point(page_w - margin - 190, 34), mgr_text, fontsize=7.5, color=(0.3, 0.35, 0.4), fontfile=font_path)

    page.draw_line(pymupdf.Point(margin, 56), pymupdf.Point(page_w - margin, 56), color=(0.8, 0.85, 0.9), width=0.8)

    # Document Title
    inv_no = data.get('invoice_no', 'PRO-2026/09-89')
    inv_date = data.get('invoice_date', '29.09.2026')
    page.insert_text(pymupdf.Point(margin, 74), f"FACTURĂ PROFORMĂ  /  PROFORMA INVOICE  № {inv_no}", fontsize=11, color=(0.11, 0.16, 0.29), fontfile=font_path)
    page.insert_text(pymupdf.Point(page_w - margin - 100, 74), f"Data: {inv_date}", fontsize=9, color=(0.4, 0.4, 0.4), fontfile=font_path)

    # Metadata Box
    page.draw_rect(pymupdf.Rect(margin, 82, page_w - margin, 126), color=(0.8, 0.85, 0.9), fill=(0.96, 0.97, 0.99))
    supplier_info = f"Furnizor: {data.get('supplier', 'Lombardo S.p.a. (Villongo, BG, Italia)')}\nConditii: {data.get('incoterms', 'EXW Villongo')} | Valuta: {data.get('currency', 'EUR')}"
    trans_info = f"Document transport: {data.get('cmr_no', 'CMR MD-IT 89412')}\nCamion / Trailer: {data.get('truck_no', 'C LK 421 / TR 890')}"
    page.insert_text(pymupdf.Point(margin + 8, 98), supplier_info, fontsize=8, color=(0.2, 0.25, 0.3), fontfile=font_path)
    page.insert_text(pymupdf.Point(page_w / 2 + 10, 98), trans_info, fontsize=8, color=(0.2, 0.25, 0.3), fontfile=font_path)

    # Table Header
    col_x = [
        margin,                  # 0: №
        margin + 20,             # 1: Cod (65)
        margin + 85,             # 2: Denumire marfa (RO / Original) (205)
        margin + 290,            # 3: HS Code (65)
        margin + 355,            # 4: Tara (25)
        margin + 380,            # 5: Cant. (30)
        margin + 410,            # 6: Masa Net (45)
        margin + 455,            # 7: Pret € (40)
        page_w - margin          # 8: Total € (40)
    ]

    header_y = 136
    row_h = 16
    page.draw_rect(pymupdf.Rect(margin, header_y, page_w - margin, header_y + row_h), color=(0.11, 0.16, 0.29), fill=(0.11, 0.16, 0.29))

    th = ["№", "Cod Articol", "Denumire Marfa (RO / Original)", "Cod Vamal", "Tara", "Cant.", "Net kg", "Pret €", "Total €"]
    for i in range(len(th)):
        page.insert_text(pymupdf.Point(col_x[i] + 2, header_y + 11), th[i], fontsize=7, color=(1, 1, 1), fontfile=font_path)

    y = header_y + row_h
    items = data.get("items", [])
    tot_qty = 0
    tot_net = 0.0
    tot_val = 0.0

    for idx, it in enumerate(items, 1):
        if y > page_h - 90:
            page = doc.new_page(width=page_w, height=page_h)
            y = 40
            page.draw_rect(pymupdf.Rect(margin, y, page_w - margin, y + row_h), color=(0.11, 0.16, 0.29), fill=(0.11, 0.16, 0.29))
            for i in range(len(th)):
                page.insert_text(pymupdf.Point(col_x[i] + 2, y + 11), th[i], fontsize=7, color=(1, 1, 1), fontfile=font_path)
            y += row_h

        if idx % 2 == 0:
            page.draw_rect(pymupdf.Rect(margin, y, page_w - margin, y + row_h), color=(0.97, 0.98, 0.99), fill=(0.97, 0.98, 0.99))

        code = it.get("code") or it.get("articol") or f"ART-{idx}"
        raw_name = it.get("name") or it.get("denumire") or "Corp iluminat"
        ro_name = translate_to_ro(raw_name)
        hs = it.get("hs_code", match_hs_code(raw_name)["code"])
        country = it.get("origin_country", "IT")
        qty = float(it.get("qty", 1))
        net = float(it.get("net_kg", round(qty * 0.85, 2)))
        price = float(it.get("price_eur", it.get("unit_price", 50.0)))
        tot = round(qty * price, 2)

        tot_qty += qty
        tot_net += net
        tot_val += tot

        name_display = f"{ro_name}"[:44]
        page.insert_text(pymupdf.Point(col_x[0] + 2, y + 11), str(idx), fontsize=7, color=(0.2, 0.2, 0.2), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[1] + 2, y + 11), str(code)[:13], fontsize=7, color=(0.1, 0.1, 0.1), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[2] + 2, y + 11), name_display, fontsize=6.8, color=(0.1, 0.1, 0.1), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[3] + 2, y + 11), hs[:11], fontsize=6.5, color=(0.3, 0.3, 0.3), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[4] + 2, y + 11), country, fontsize=7, color=(0.3, 0.3, 0.3), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[5] + 2, y + 11), f"{qty:.0f}" if qty.is_integer() else f"{qty:.1f}", fontsize=7, color=(0.2, 0.2, 0.2), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[6] + 2, y + 11), f"{net:.1f}", fontsize=7, color=(0.2, 0.2, 0.2), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[7] + 2, y + 11), f"{price:.1f}", fontsize=7, color=(0.2, 0.2, 0.2), fontfile=font_path)
        page.insert_text(pymupdf.Point(col_x[8] - 35, y + 11), f"{tot:.1f}", fontsize=7, color=(0.1, 0.1, 0.1), fontfile=font_path)

        y += row_h

    # Total Row
    page.draw_rect(pymupdf.Rect(margin, y, page_w - margin, y + row_h), color=(0.85, 0.9, 0.95), fill=(0.85, 0.9, 0.95))
    page.insert_text(pymupdf.Point(col_x[1] + 2, y + 11), "TOTAL GENERAL MARFA:", fontsize=7.5, color=(0.1, 0.15, 0.25), fontfile=font_path)
    page.insert_text(pymupdf.Point(col_x[5] + 2, y + 11), f"{tot_qty:.0f}", fontsize=7.5, color=(0.1, 0.15, 0.25), fontfile=font_path)
    page.insert_text(pymupdf.Point(col_x[6] + 2, y + 11), f"{tot_net:.1f} kg", fontsize=7.5, color=(0.1, 0.15, 0.25), fontfile=font_path)
    page.insert_text(pymupdf.Point(col_x[8] - 45, y + 11), f"€ {tot_val:,.2f}", fontsize=8, color=(0.11, 0.16, 0.29), fontfile=font_path)
    y += row_h + 15

    # Legal Declaration & Signatures
    decl = (
        "Declaratie pe proprie raspundere: Preturile mentionate reflecta valoarea reala a tranzactiei pentru vamuire conform legilor RM.\n"
        "Marfurile sunt certificate CE. Fara TVA conform regimului vamal de tranzit."
    )
    page.insert_text(pymupdf.Point(margin, y), decl, fontsize=6.5, color=(0.4, 0.4, 0.4), fontfile=font_path)
    y += 28

    page.draw_line(pymupdf.Point(margin, y), pymupdf.Point(margin + 180, y), color=(0.6, 0.6, 0.6), width=0.5)
    page.draw_line(pymupdf.Point(page_w - margin - 180, y), pymupdf.Point(page_w - margin, y), color=(0.6, 0.6, 0.6), width=0.5)
    page.insert_text(pymupdf.Point(margin, y + 10), "Expeditor / Stampila furnizor", fontsize=7, color=(0.3, 0.3, 0.3), fontfile=font_path)
    page.insert_text(pymupdf.Point(page_w - margin - 180, y + 10), "Importator / Director Smart Light SRL", fontsize=7, color=(0.3, 0.3, 0.3), fontfile=font_path)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path

def main():
    parser = argparse.ArgumentParser(description="Moldovan Customs Translator & Proforma Generator")
    parser.add_argument("--input", required=True, help="Input BOM or Invoice JSON file")
    parser.add_argument("--out-xlsx", default="Specificatie_Vamala.xlsx", help="Output customs Excel specification")
    parser.add_argument("--out-pdf", default="Factura_Proforma.pdf", help="Output customs PDF proforma invoice")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    # If it's BOM_v3.json format
    if "items" not in data:
        if "master_bom_items" in data:
            data["items"] = data["master_bom_items"]
        elif "full_bom" in data:
            data["items"] = data["full_bom"]
        elif "equipment" in data:
            data["items"] = data["equipment"]
        elif "bom" in data:
            data["items"] = data["bom"]

    # Normalize fields
    for it in data.get("items", []):
        if "code" not in it and "article" in it:
            it["code"] = it["article"]
        if "unit_price" not in it and "price_eur" in it:
            it["unit_price"] = it["price_eur"]

    print(f"[*] Processing {len(data.get('items', []))} items for Moldovan customs...")
    xlsx_res = generate_customs_excel(data, args.out_xlsx)
    print(f"[+] Customs Excel Specification generated: {xlsx_res}")
    pdf_res = generate_proforma_pdf(data, args.out_pdf)
    print(f"[+] Customs Proforma PDF generated: {pdf_res}")

if __name__ == "__main__":
    main()
